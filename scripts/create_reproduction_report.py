from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(r'E:/xjtu/os/科研实习/BotRGCN_reproduction_report.docx')

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)
def borders(table):
    tblPr=table._tbl.tblPr; b=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); b.append(e)
    tblPr.append(b)
def table(doc, headers, rows):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'; borders(t)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,'1F4E78'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(9)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=str(v); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2==1: shade(cells[i],'F2F6FA')
            for p in cells[i].paragraphs:
                for r in p.runs: r.font.size=Pt(9)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t
def para(doc, text, boldlead=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.15
    if boldlead and text.startswith(boldlead): p.add_run(boldlead).bold=True; p.add_run(text[len(boldlead):])
    else: p.add_run(text)
    return p

doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
styles=doc.styles; styles['Normal'].font.name='Microsoft YaHei'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); styles['Normal'].font.size=Pt(10.5)
for name,size in [('Title',22),('Heading 1',15),('Heading 2',12)]:
    styles[name].font.name='Microsoft YaHei'; styles[name]._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); styles[name].font.size=Pt(size); styles[name].font.bold=True; styles[name].font.color.rgb=RGBColor(0,0,0)

p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('BotRGCN Twitter Bot Detection 复现论文报告')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('基于官方 TwiBot-20 数据集的完整复现与实验分析').italic=True
para(doc,'报告目的：复现 Feng 等人提出的 BotRGCN 方法，完成数据构建、论文特征编码、关系图卷积训练、主实验、特征消融和基线对比，并如实分析复现结果与原论文结果的差异。')

doc.add_heading('一 复现结论',1)
para(doc,'本项目已完成可验证的端到端复现流程。官方 TwiBot-20 数据被构造成包含 229,580 个用户节点和 227,979 条 following/follower 边的异构图；每个节点包含用户 description、tweets、6 个数值属性和 11 个类别属性。修正后的模型使用论文描述的 RoBERTa、多模态投影、两层 R-GCN、R-GCN 后 MLP 和二元交叉熵目标。')
para(doc,'在当前可复现配置下，BotRGCN 三个随机种子的平均 F1 为 0.8537，MCC 为 0.6633。论文报告的 F1 为 0.8707、MCC 为 0.7021。当前结果低于论文，但数据规模、图规模、特征类别和核心网络结构已对齐；由于论文正文未完整公开隐藏维度、学习率、训练轮数和早停等参数，本文将这些设置明确标记为复现配置，而不冒充原始超参数。')

doc.add_heading('二 论文方法与实现',1)
doc.add_heading('2.1 图数据',2)
para(doc,'节点对应 Twitter 用户。following 与 follower 被建模为两种关系类型，关系类型分别映射为 0 和 1。节点标签采用官方数据集标签，0 表示 human，1 表示 bot；训练、验证和测试划分沿用 TwiBot-20 官方划分。')
table(doc,['对象','数量/形状','处理状态'],[['节点','229,580','已验证'],['边','227,979','已验证'],['following / follower','117,110 / 110,869','已构造'],['train / validation / test','8,278 / 2,365 / 1,183','已对齐'],['support 节点','217,754','无标签，仅参与图传播']])
doc.add_heading('2.2 节点特征',2)
para(doc,'用户 description 使用 roberta-base 的首 token 表示；用户 tweets 逐条使用 RoBERTa 编码，并对用户的全部可用 tweet 表示取平均。没有可用 tweet 的用户使用零向量，这对应原始记录为空，而非编码失败。')
para(doc,'数值属性包括 followers、followings、favorites、statuses、active_days 和 screen_name_length，并使用训练集统计量进行 z-score 标准化。类别属性包括 protected、geo_enabled、verified、contributors_enabled、is_translator、is_translation_enabled、profile_background_tile、profile_user_background_image、has_extended_profile、default_profile 和 default_profile_image。类别字段按二元 one-hot 值编码后输入全连接层。')
table(doc,['特征对象','形状'],[['结构化属性','229,580 × 17'],['description RoBERTa','229,580 × 768'],['tweet RoBERTa','229,580 × 768'],['标签','229,580'],['edge_index','2 × 227,979'],['edge_type','227,979']])
doc.add_heading('2.3 模型与损失',2)
para(doc,'四类输入分别投影到 D/4 维并拼接，经过 W1 和 Leaky-ReLU 得到初始节点表示。随后使用两层 R-GCN 区分 following 和 follower 关系；最后经过 W2 和 Leaky-ReLU，再通过单 logit 分类层输出 Bot 概率。训练采用论文给出的二元交叉熵目标，并以 Adam 的 weight decay 实现 L2 正则化。')
para(doc,'当前复现配置为 D=64、学习率 1e-3、weight decay 5e-4、最多 100 个 epoch、验证集 F1 patience=20。论文正文未给出这些数值的完整明细，因此报告中将其标注为 reproduction configuration。')

doc.add_heading('三 实验环境与流程',1)
para(doc,'实验在 Windows、Python 3.13、PyTorch、PyTorch Geometric、Transformers 和 NVIDIA RTX 4060 Laptop GPU 上完成。RoBERTa 文本表示先离线编码并与全量节点 ID 对齐，再执行全图节点分类训练。所有正式结果均使用固定随机种子重复三次，报告均值和样本标准差。')
para(doc,'正式数据检查确认所有特征有限、节点和边索引一致、训练/验证/测试 mask 无交集且覆盖全部节点。support 节点没有标签，但参与图消息传播。')

doc.add_heading('四 主实验结果',1)
table(doc,['指标','均值 ± 标准差'],[['Accuracy','0.8166 ± 0.0009'],['Precision','0.7509 ± 0.0014'],['Recall','0.9891 ± 0.0054'],['F1','0.8537 ± 0.0012'],['MCC','0.6633 ± 0.0042']])
table(doc,['指标','论文结果','本复现','差值'],[['Accuracy','0.8462','0.8166','-0.0296'],['F1','0.8707','0.8537','-0.0170'],['MCC','0.7021','0.6633','-0.0388']])
para(doc,'本复现的 F1 和 MCC 低于论文报告值，但三次运行的 F1 标准差为 0.0012，说明当前实现具有稳定的重复性。差异可能来自论文未公开的隐藏维度、优化器参数、训练策略和具体 RoBERTa 处理细节；本文不将这些未公开因素解释为确定性原因。')

doc.add_heading('五 特征消融实验',1)
para(doc,'下表使用修正后的 BCE 版本，每种配置运行三个随机种子。早期简化特征实验和旧版 CrossEntropy 结果不纳入本报告。')
table(doc,['设置','Accuracy','F1','MCC'],[['完整特征','0.8155 ± 0.0010','0.8528 ± 0.0012','0.6611 ± 0.0040'],['去除 description','0.8163 ± 0.0022','0.8536 ± 0.0018','0.6633 ± 0.0054'],['去除 tweets','0.8177 ± 0.0013','0.8544 ± 0.0006','0.6652 ± 0.0021'],['去除数值属性','0.8157 ± 0.0009','0.8533 ± 0.0009','0.6627 ± 0.0033'],['去除类别属性','0.6915 ± 0.0400','0.7549 ± 0.0142','0.3885 ± 0.0669'],['仅保留属性','0.6205 ± 0.0688','0.6842 ± 0.0155','0.2137 ± 0.1851']])
para(doc,'类别属性移除后性能明显下降，是当前实验中最清晰的特征贡献。去除 description、tweets 或数值属性后性能没有下降，反而略有上升，因此不能在本复现中宣称每类特征都带来独立增益。')

doc.add_heading('六 关系消融实验',1)
table(doc,['关系设置','Accuracy','F1','MCC'],[['all','0.8163 ± 0.0005','0.8536 ± 0.0005','0.6633 ± 0.0016'],['following','0.7681 ± 0.0788','0.8169 ± 0.0605','0.5590 ± 0.1733'],['follower','0.8163 ± 0.0022','0.8541 ± 0.0017','0.6654 ± 0.0049']])
para(doc,'following 的 seed=44 结果明显低于另两次，导致该设置方差较大；因此不能仅依据三次均值断言 all 一定优于单一关系。正式报告应保留每次运行结果，并将该不稳定性作为限制。')

doc.add_heading('七 基线实验',1)
para(doc,'基线使用相同的四类多模态特征投影、BCE 目标和三次随机种子。MLP 不使用图边；GCN 和 GAT 使用同样的两层图卷积对照。它们是本项目统一框架下的实现，用于相对比较，不等同于论文未完整公开细节的作者基线代码。')
table(doc,['模型','Accuracy','F1','MCC'],[['MLP','0.8169 ± 0.0005','0.8550 ± 0.0003','0.6688 ± 0.0010'],['GCN','0.5900 ± 0.0301','0.7139 ± 0.0092','0.1831 ± 0.0743'],['GAT','0.5424 ± 0.0024','0.7026 ± 0.0009','0.0205 ± 0.0356'],['BotRGCN','0.8166 ± 0.0009','0.8537 ± 0.0012','0.6633 ± 0.0042']])
para(doc,'在当前配置下，MLP 略高于 BotRGCN；GCN 和 GAT 明显较低。该结果与论文“R-GCN 优于其他结构”的定性结论并不完全一致，可能反映了尚未公开的基线超参数和实现细节差异。报告保留这一结果，以保证复现过程可审计。')

doc.add_heading('八 复现局限与后续工作',1)
para(doc,'第一，论文正文没有完整公开 D、学习率、训练轮数、早停和 dropout 等训练参数，因此当前实验无法声称逐参数复刻作者训练。第二，RoBERTa 的具体 tokenizer 截断长度、空 description 处理和作者实现细节未完全公开。第三，三次种子足以进行稳定性初检，但仍少于更严格实验通常采用的五次或更多重复。第四，当前 GCN/GAT/MLP 是统一框架下的对照实现，不能替代作者原始基线代码。')
para(doc,'后续若取得作者代码或补充材料，应优先核对隐藏维度、RoBERTa 预处理、优化器和训练轮数，并在同一数据和随机种子协议下重跑主实验与消融。')

doc.add_heading('九 项目产物',1)
for x in ['scripts/train_botrgcn_official.py：正式 BotRGCN 训练','scripts/train_botrgcn_ablation.py：特征与关系消融','scripts/train_baselines.py：MLP、GCN、GAT 基线','scripts/validate_full_features.py：全量数据对齐检查','processed/：节点、边、属性、RoBERTa 特征和 mask','reproduction_status.md：实验状态与结果记录']:
    doc.add_paragraph(x, style='List Bullet')
para(doc,'结论：本项目已经形成一套可运行、可检查、结果可重复的 BotRGCN 复现实现，适合作为课程/科研实习中的复现项目展示。向老师提交时，应将本文定位为“基于公开论文和官方 TwiBot-20 数据的独立复现与差异分析”，而不是声称完全恢复了作者未公开的全部训练细节。')

doc.save(OUT); print(OUT)
