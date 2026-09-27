from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path(r'E:/xjtu/os/科研实习/BotRGCN_文件附录.docx')
def shade(c,fill):
    shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); c._tc.get_or_add_tcPr().append(shd)
def borders(t):
    b=OxmlElement('w:tblBorders')
    for x in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+x); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); b.append(e)
    t._tbl.tblPr.append(b)
def table(doc,heads,rows):
    t=doc.add_table(rows=1,cols=len(heads)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER; borders(t)
    for i,h in enumerate(heads):
        c=t.rows[0].cells[i]; c.text=h; shade(c,'1F4E78'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(9)
    for n,row in enumerate(rows):
        cs=t.add_row().cells
        for i,v in enumerate(row):
            cs[i].text=str(v); cs[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if n%2: shade(cs[i],'F2F6FA')
            for p in cs[i].paragraphs:
                for r in p.runs: r.font.size=Pt(9)
    doc.add_paragraph()
def p(doc,s):
    q=doc.add_paragraph(s); q.paragraph_format.space_after=Pt(5); q.paragraph_format.line_spacing=1.12
    return q
doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
for st,size in [('Normal',10.5),('Title',21),('Heading 1',15),('Heading 2',12)]:
    doc.styles[st].font.name='Microsoft YaHei'; doc.styles[st]._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); doc.styles[st].font.size=Pt(size); doc.styles[st].font.color.rgb=RGBColor(0,0,0)
title=doc.add_paragraph(style='Title'); title.alignment=WD_ALIGN_PARAGRAPH.CENTER; title.add_run('BotRGCN 复现项目文件附录')
q=doc.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.add_run('与复现论文报告配套提交').italic=True
p(doc,'本附录说明复现项目的代码组织、数据产物、运行入口和提交建议，便于老师复核实验过程并在另一台机器上重现实验。附录不包含 TwiBot-20 原始大文件和 Hugging Face 模型缓存；这些文件应按照数据许可和存储条件单独提供。')
doc.add_heading('一 建议提交目录',1)
table(doc,['目录/文件','用途','提交建议'],[
['BotRGCN_reproduction_report.docx','复现论文报告','提交'],['BotRGCN_文件附录.docx','本附录','提交'],['botrgcn-reproduction/scripts/','全部处理、编码、训练脚本','提交'],['botrgcn-reproduction/reproduction_status.md','过程状态和结果记录','提交'],['botrgcn-reproduction/README.md','项目说明','提交并补充'],['botrgcn-reproduction/requirements-base.txt','基础依赖记录','提交'],['botrgcn-reproduction/environment.md','环境记录','提交'],['botrgcn-reproduction/processed/','特征、边、标签和 mask','可选择压缩提交'],['botrgcn-reproduction/Twibot-20/','官方原始数据','不直接放入邮件/仓库，按许可单独提供']])
doc.add_heading('二 核心代码清单',1)
table(doc,['文件','功能','复核重点'],[
['scripts/build_node_index.py','构造全量节点 ID 到连续索引','节点数 229580'],['scripts/build_full_edges.py','扫描 train/dev/test/support 并构图','边数 227979，关系类型 2'],['scripts/build_paper_property_features.py','构造 6 个数值属性和 11 个类别属性','训练集 z-score'],['scripts/encode_labeled_roberta.py','编码有标签用户文本','description/tweets 768 维'],['scripts/encode_support_roberta.py','分块编码 support 文本','支持断点续跑'],['scripts/merge_support_roberta_chunks.py','合并 support chunk','与全量节点对齐'],['scripts/validate_full_features.py','验证形状、有限值和 mask','正式训练前运行'],['scripts/train_botrgcn_official.py','正式 BotRGCN 主实验','BCE、R-GCN、MCC'],['scripts/train_botrgcn_ablation.py','特征与关系消融','BCE 版本'],['scripts/train_baselines.py','MLP、GCN、GAT 基线','统一多模态输入']])
doc.add_heading('三 关键数据产物',1)
table(doc,['文件','内容','形状/数量'],[
['processed/full_node_index.json','节点 ID 映射','229580 节点'],['processed/edge_index.pt','边索引','2 × 227979'],['processed/edge_type.pt','关系类型','227979'],['processed/x_paper_properties.pt','论文属性特征','229580 × 17'],['processed/description_roberta_full.pt','description 表示','229580 × 768'],['processed/tweet_roberta_full.pt','tweet 平均表示','229580 × 768'],['processed/y_full.pt','标签','229580'],['processed/train_mask.pt','训练 mask','8278'],['processed/val_mask.pt','验证 mask','2365'],['processed/test_mask.pt','测试 mask','1183']])
doc.add_heading('四 推荐运行顺序',1)
for x in ['建立虚拟环境并安装 requirements-base.txt、PyTorch、PyTorch Geometric 和 Transformers。','准备 TwiBot-20 原始文件 train.json、dev.json、test.json、support.json。','运行 build_node_index.py 和 build_full_edges.py。','运行 build_paper_property_features.py。','运行 encode_labeled_roberta.py；support 编码使用 encode_support_roberta.py 并可分块断点续跑。','运行 merge_support_roberta_chunks.py 和 validate_full_features.py。','运行 train_botrgcn_official.py --seed 42、--seed 43、--seed 44。','运行 train_botrgcn_ablation.py 和 train_baselines.py，记录 Accuracy、F1、MCC。']:
    doc.add_paragraph(x,style='List Number')
doc.add_heading('五 关键命令',1)
for x in ['python scripts/validate_full_features.py','python scripts/train_botrgcn_official.py --seed 42','python scripts/train_botrgcn_official.py --seed 43','python scripts/train_botrgcn_official.py --seed 44','python scripts/train_botrgcn_ablation.py --ablation no-categorical --relation all --seed 42','python scripts/train_botrgcn_ablation.py --ablation full --relation following --seed 42','python scripts/train_baselines.py --kind mlp --seed 42','python scripts/train_baselines.py --kind gcn --seed 42','python scripts/train_baselines.py --kind gat --seed 42']:
    doc.add_paragraph(x,style='List Bullet')
doc.add_heading('六 最终结果摘要',1)
table(doc,['实验','Accuracy','F1','MCC'],[['BotRGCN 主实验','0.8166 ± 0.0009','0.8537 ± 0.0012','0.6633 ± 0.0042'],['MLP 基线','0.8169 ± 0.0005','0.8550 ± 0.0003','0.6688 ± 0.0010'],['GCN 基线','0.5900 ± 0.0301','0.7139 ± 0.0092','0.1831 ± 0.0743'],['GAT 基线','0.5424 ± 0.0024','0.7026 ± 0.0009','0.0205 ± 0.0356']])
p(doc,'结果解释：主实验和基线均采用当前统一的多模态特征与 BCE 配置。MLP 略高于 BotRGCN 的结果应如实保留；这说明当前复现结果与论文中 R-GCN 优于其他结构的定性结论存在差异。')
doc.add_heading('七 不建议提交的文件',1)
for x in ['Twibot-20/support.json 等原始大文件：体积大且应遵守数据许可。','hf_cache/：RoBERTa 模型缓存，体积大，不应放入代码仓库。','support_roberta_chunks_backup/：仅用于故障恢复的临时备份。','旧版 CE 实验输出和简化 profile/behavior 实验：不属于最终复现结果。','Python __pycache__/、临时日志和渲染中间文件。']:
    doc.add_paragraph(x,style='List Bullet')
doc.add_heading('八 提交前检查',1)
for x in ['报告中的主实验结果与本附录一致。','所有最终结果均来自 BCE 修正版脚本。','原始数据、RoBERTa 缓存和临时备份未混入代码压缩包。','requirements、GPU/CUDA、Python 版本已记录。','老师可根据运行顺序找到正式训练入口。']:
    doc.add_paragraph('□ '+x)
p(doc,'项目定位：公开论文和官方 TwiBot-20 数据基础上的独立复现、工程实现和差异分析。')
doc.save(OUT); print(OUT)
